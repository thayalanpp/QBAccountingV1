"""
split_reader.py

Reads the monthly split workbook (e.g. C:\\NM\\2025yearend\\NM_2025.xlsx) exactly
as the bookkeeper left it - the year is closed, so these are the numbers of
record; differences are reported as adjustments, never "fixed" here.

One sheet per month ("Sep 2024", "July 2025", ...). Only the split block is used:

    row   A (category)   B (total)            C (DPC portion)                 D
    2     Salary         hygiene salaries     B*0.6
    3     rent           3,225.38             B*0.6
    4     lease          125.43 / 352.54      B*0.6 or B*0.8
    5     bookkeping     250                  B*0.5
    6     visa/expense   visa payment due     (B6-D6)*0.8 + D6               D6 = lab fees on the card
    8                    SUM(B2:B6)           C8 = SUM(C2:C6) - D8  = the DNTL split cheque

    Lab fees on the card: descriptions in column K, amounts in column M (rows 1-7), M8 = total.

Values are read from Excel's saved results (the file must have been saved in
Excel). The formula text is read too, so the rate each month used is recorded.
"""

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Optional

from openpyxl import load_workbook

CATEGORY_ROWS = range(2, 7)          # rows 2-6: salary, rent, lease, bookkeeping, visa
USUAL_RATES = {"salary": [0.6], "rent": [0.6], "lease": [0.6, 0.8], "bookkeeping": [0.5]}
VISA_DPC_RATE = 0.8                  # non-lab card charges; lab fees are 100% DPC


def _category(label: str) -> str:
    t = re.sub(r"[^a-z]", "", str(label or "").lower())
    if "visa" in t:                                   # "Visa", "visa payment", "VISA due" ...
        return "visa"
    return {"bookkeping": "bookkeeping"}.get(t, t)


def _num(v) -> float:
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):                            # a number typed as text, e.g. "10,076.74" or "$1,234.50"
        t = v.strip().replace(",", "").replace("$", "")
        try:
            return float(t)
        except ValueError:
            return 0.0
    return 0.0


@dataclass
class SplitLine:
    category: str
    label: str
    total: float                 # column B
    dpc_portion: float           # column C, as is
    formula: Optional[str]       # column C formula text, e.g. "=+B4*0.8"

    @property
    def rate(self) -> Optional[float]:
        return round(self.dpc_portion / self.total, 4) if self.total else None


@dataclass
class LabItem:
    description: str
    amount: float
    row: int


@dataclass
class SplitMonth:
    sheet: str
    month: date                  # first day of the sheet's month
    lines: list[SplitLine]
    lab_fees: float              # D6
    lab_items: list[LabItem]
    cheque_total: float          # C8
    other_deduction: float       # D8
    notes: list[str] = field(default_factory=list)

    def line(self, category: str) -> Optional[SplitLine]:
        return next((l for l in self.lines if l.category == category), None)

    @property
    def visa_due(self) -> float:
        v = self.line("visa")
        return v.total if v else 0.0

    @property
    def visa_dpc(self) -> float:
        v = self.line("visa")
        return v.dpc_portion if v else 0.0

    @property
    def visa_hygiene(self) -> float:
        return round(self.visa_due - self.visa_dpc, 2)


def _sheet_month(name: str) -> Optional[date]:
    for fmt in ("%b %Y", "%B %Y"):
        try:
            d = datetime.strptime(name.strip(), fmt)
            return date(d.year, d.month, 1)
        except ValueError:
            pass
    return None


def read_split_workbook(path: str) -> list[SplitMonth]:
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    months = []
    for ws in values.worksheets:
        month = _sheet_month(ws.title)
        if month is None:
            continue                                  # not a month sheet
        wf = formulas[ws.title]
        notes = []
        lines = []
        for r in CATEGORY_ROWS:
            label = ws.cell(r, 1).value
            if not label:
                continue
            b, c = ws.cell(r, 2).value, ws.cell(r, 3).value
            for col, v in ((2, b), (3, c)):
                if isinstance(wf.cell(r, col).value, str) and wf.cell(r, col).value.startswith("=") and v is None:
                    notes.append(f"row {r}: column {'BC'[col - 2]} has a formula ({wf.cell(r, col).value}) but no saved "
                                 f"value - open and save the file in Excel")
            if isinstance(b, str):
                notes.append(f"row {r}: column B is text ({b!r}), read as {_num(b):,.2f}")
            f = wf.cell(r, 3).value
            cat = _category(label)
            if r == 6 and cat not in ("salary", "rent", "lease", "bookkeeping"):
                cat = "visa"                          # row 6 is the card line whatever it's called ("visa", "expense", ...)
            lines.append(SplitLine(cat, str(label).strip(), round(_num(b), 2), _num(c),
                                   f if isinstance(f, str) else None))
        lab_items = []
        for r in range(1, 8):
            amt = ws.cell(r, 13).value                 # M
            if isinstance(amt, (int, float)) and amt:
                lab_items.append(LabItem(str(ws.cell(r, 11).value or "").strip(), round(float(amt), 2), r))
        months.append(SplitMonth(sheet=ws.title, month=month, lines=lines, lab_fees=round(_num(ws["D6"].value), 2),
                                 lab_items=lab_items, cheque_total=_num(ws["C8"].value),
                                 other_deduction=_num(ws["D8"].value), notes=notes))
    return sorted(months, key=lambda m: m.month)
