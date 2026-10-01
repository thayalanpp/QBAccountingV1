"""
split_pipeline.py  -  checks the monthly Dental / Hygiene split workbook

    read_workbook ─▶ load_evidence ─▶ assess ─▶ rollforward ─▶ save_to_sql ─▶ write_report ─▶ END

    read_workbook   split_reader.py - NM_<year>.xlsx, as is (the year is closed)
    load_evidence   card statements (visa_pipeline), Dental + Hygiene banks
                    (bank_pipeline), Dental QuickBooks (run_qb_load)
    assess          split_assess.py - per month: workbook, links, cash test,
                    expense test, QuickBooks postings
    rollforward     rollforward.py - QB 22200 vs the real card balance at both year-ends,
                    itemized; and the year-end cut-off of Hygiene's share
    save_to_sql     fin.SplitMonths / SplitLines / SplitLabItems + fin.Findings
    write_report    <year>yearend\\split_review_FY<year>.xlsx - one line per month,
                    proposed adjustments, lab fees, QB postings, findings

Run it after the other pipelines have loaded the year:
    python split_pipeline.py
    python split_pipeline.py --workbook "C:\\NM\\2025yearend\\NM_2025.xlsx"

Nothing in QuickBooks or the workbook is changed.
"""

import argparse
import os
import sys
from datetime import timedelta
from typing import Optional, TypedDict

from langgraph.graph import END, StateGraph
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import split_sql
from qb_loader import fiscal_period
import rollforward as rfmod
from split_assess import Evidence, MonthResult, assess_month
from split_reader import read_split_workbook
from sql_helper import ensure_schema, get_engine

# ---------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------
FISCAL_YEAR_END = 2025
YEAR_ROOT = r"C:\NM\{year}yearend"
WORKBOOK = r"NM_{year}.xlsx"                       # in YEAR_ROOT
REPORT = r"split_review_FY{year}.xlsx"             # written to YEAR_ROOT

CARD_ACCOUNT = "TD Aeroplan Visa Infinite Privilege"          # fin.Accounts (visa_pipeline)
DNTL_BANK = "Scotiabank Business 26492-00281-18"              # fin.Accounts (bank_pipeline)
HYG_BANK = "Scotiabank Business 26492-00688-10"
QB_COMPANY = "dntl"
DNTL_BANK_QB_ACCOUNT = "10000"
CARD_QB_ACCOUNT = "22200"
VISA_PAYMENT_PATTERN = "td visa"                  # Hygiene's payments to TD look like "Pc-Td Visa 38099194"

ENGINE = None


class SplitState(TypedDict):
    workbook: str
    report: str
    fy: object
    months: list
    boundary: list                  # the month before and after the year (from the neighbouring workbooks)
    evidence: Optional[Evidence]
    results: list
    boundary_results: list
    rollforward: object
    cutoffs: list
    findings: list
    error: Optional[str]


def _add_months(d, n):
    y, m = divmod(d.year * 12 + d.month - 1 + n, 12)
    return d.replace(year=y, month=m + 1, day=1)


def _neighbour_workbook(workbook: str, fy_end: int, year: int) -> Optional[str]:
    """C:\\NM\\2025yearend\\NM_2025.xlsx -> C:\\NM\\2024yearend\\NM_2024.xlsx for year 2024."""
    path = workbook.replace(f"{fy_end}yearend", f"{year}yearend").replace(f"NM_{fy_end}", f"NM_{year}")
    return path if path != workbook and os.path.exists(path) else None


def read_workbook(state: SplitState) -> dict:
    fy = state["fy"]
    try:
        all_months = read_split_workbook(state["workbook"])
    except Exception as e:
        return {"error": f"can't read {state['workbook']}: {e}"}
    months = [m for m in all_months if fy.start <= m.month <= fy.end]
    print(f"📘 read_workbook: {len(months)} month sheets in FY{fy.fy_end_year}: " + ", ".join(m.sheet for m in months))

    # Boundary months: the last month of the previous year and the first of the next. They're
    # linked and checked like any month but kept out of this year's totals.
    boundary = []
    for target, year in ((_add_months(fy.start, -1), fy.fy_end_year - 1), (_add_months(fy.end, 1), fy.fy_end_year + 1)):
        found = next((m for m in all_months if m.month == target), None)
        src = state["workbook"] if found else None
        if not found:
            other = _neighbour_workbook(state["workbook"], fy.fy_end_year, year)
            if other:
                found = next((m for m in read_split_workbook(other) if m.month == target), None)
                src = other if found else None
        if found:
            boundary.append(found)
            print(f"   boundary month {found.sheet} from {os.path.basename(src)}")
        else:
            print(f"   (no {target:%b %Y} sheet found - boundary month skipped)")
    return {"months": months, "boundary": boundary, "error": None if months else "no month sheets inside the fiscal year"}


def load_evidence(state: SplitState) -> dict:
    fy = state["fy"]
    ev = Evidence(cycles=split_sql.load_card_cycles(ENGINE, CARD_ACCOUNT),
                  dntl_bank=split_sql.load_bank_lines(ENGINE, DNTL_BANK),
                  hyg_bank=split_sql.load_bank_lines(ENGINE, HYG_BANK),
                  all_withdrawals=split_sql.load_all_bank_withdrawals(ENGINE),
                  qb_dntl=split_sql.load_qb_lines(ENGINE, QB_COMPANY, fy.start - timedelta(days=60),
                                                  fy.end + timedelta(days=90)),
                  dntl_bank_qb_account=DNTL_BANK_QB_ACCOUNT, card_qb_account=CARD_QB_ACCOUNT,
                  visa_payment_pattern=VISA_PAYMENT_PATTERN)
    print(f"🔎 load_evidence: {len(ev.cycles)} card statements, {len(ev.dntl_bank)} Dental bank lines, "
          f"{len(ev.hyg_bank)} Hygiene bank lines, {len(ev.qb_dntl)} Dental QB lines")
    missing = [n for n, v in [("card statements", ev.cycles), ("Dental bank", ev.dntl_bank),
                              ("Hygiene bank", ev.hyg_bank), ("Dental QuickBooks", ev.qb_dntl)] if not v]
    if missing:
        print(f"   ⚠️ not loaded yet: {', '.join(missing)} - those checks will be skipped")
    if not ev.hyg_bank:
        print("   ⚠️ the TD Visa is paid from Hygiene's bank - without it the cash and expense tests can't run. "
              "Load it first:  python bank_pipeline.py --only hyg")
    return {"evidence": ev}


def _print_result(r, tag=""):
    m = r.month
    print(f"   {m.sheet:10} cheque {r.cheque_number or '-':>4}  B6 {m.visa_due:>10,.2f}  "
          f"paid {r.bank_paid if r.bank_paid is not None else '-':>10}  base "
          f"{r.expense_base if r.expense_base is not None else '-':>10}  -> {r.status}{tag}")
    for f in r.findings:
        if f["severity"] == "Issue":
            print(f"      ❗ {f['area']}: {f['message']}")


def assess(state: SplitState) -> dict:
    fy, ev = state["fy"], state["evidence"]
    before = [m for m in state["boundary"] if m.month < fy.start]
    after = [m for m in state["boundary"] if m.month > fy.end]
    results, boundary_results, lab_claims = [], [], {}
    prev = None
    for m in before:                       # first, so its lab fees are claimed before this year's
        r = assess_month(m, None, ev, lab_claims)
        boundary_results.append(r)
        prev = m
        _print_result(r, "  (boundary - not in totals)")
    for m in state["months"]:
        r = assess_month(m, prev, ev, lab_claims)
        results.append(r)
        prev = m
        _print_result(r)
    for m in after:
        r = assess_month(m, prev, ev, lab_claims)
        boundary_results.append(r)
        _print_result(r, "  (boundary - not in totals)")
    findings = [f for r in results for f in r.findings]
    return {"results": results, "boundary_results": boundary_results, "findings": findings}


def rollforward(state: SplitState) -> dict:
    """22200 against the real card balance at both year-ends, and the year-end cut-off."""
    fy, ev = state["fy"], state["evidence"]
    qb_22200 = [q for q in ev.qb_dntl if q.account_number == CARD_QB_ACCOUNT]
    balances = split_sql.load_account_balance(ENGINE, QB_COMPANY, CARD_QB_ACCOUNT, fy.fy_end_year)
    matches = split_sql.load_card_matches(ENGINE, CARD_ACCOUNT)
    rf = rfmod.build_rollforward(fy.start, fy.end, ev.cycles, qb_22200, balances, matches)
    cutoffs = rfmod.build_cutoffs(ev.cycles, fy.start, fy.end,
                                  rfmod.lab_keywords(state["months"] + state["boundary"]))
    findings = list(state["findings"])
    subj = f"22200 roll-forward FY{fy.fy_end_year}"
    for n in rf.notes:
        findings.append({"subject": subj, "area": "rollforward", "severity": "Warning", "message": n})
    if rf.diff_open is not None and rf.diff_close is not None:
        print(f"📒 rollforward 22200: QB owed {rf.qb_open:,.2f} vs card {rf.card_open:,.2f} on {rf.open_date} "
              f"(diff {rf.diff_open:+,.2f})  ->  QB owed {rf.qb_close:,.2f} vs card {rf.card_close:,.2f} on "
              f"{rf.close_date} (diff {rf.diff_close:+,.2f})")
        for g, total, n in rf.groups():
            print(f"      {g:70} {total:>12,.2f}  ({n} lines)")
        print(f"      {'unexplained':70} {rf.unexplained:>12,.2f}")
        findings.append({"subject": subj, "area": "rollforward", "severity": "Issue" if abs(rf.unexplained) >= 0.01 else "Info",
                         "message": f"22200 differs from the card by {rf.diff_open:+,.2f} on {rf.open_date} and "
                                    f"{rf.diff_close:+,.2f} on {rf.close_date}; itemized {rf.explained:+,.2f}, "
                                    f"unexplained {rf.unexplained:+,.2f}", "amount": rf.diff_close})
    else:
        print("📒 rollforward 22200: skipped - " + "; ".join(rf.notes))
    for co in cutoffs:
        print(f"✂️  cut-off {co.label}: charges {co.charges:,.2f}, lab {co.lab:,.2f} -> Hygiene share {co.hygiene_share:,.2f}")
        findings.append({"subject": f"Cut-off FY{fy.fy_end_year}", "area": "cutoff", "severity": "Info",
                         "message": f"{co.label} ({co.window[0]}..{co.window[1]}, statement {co.statement_date}): "
                                    f"charges {co.charges:,.2f}, lab {co.lab:,.2f}, Hygiene share {co.hygiene_share:,.2f}",
                         "amount": co.hygiene_share})
    return {"rollforward": rf, "cutoffs": cutoffs, "findings": findings}


def save_to_sql(state: SplitState) -> dict:
    split_sql.save_results(ENGINE, state["fy"].fy_end_year, os.path.basename(state["workbook"]),
                           state["results"], state["findings"])
    print(f"💾 save_to_sql: {len(state['results'])} months, {len(state['findings'])} findings (fin.Findings, Source='split')")
    return {}


# ---------------------------------------------------------------------
# Excel review report
# ---------------------------------------------------------------------
F, FB = Font(name="Arial"), Font(name="Arial", bold=True)
HDR = Font(name="Arial", bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F4E78")
ISSUE = PatternFill("solid", fgColor="FCE4D6")
NUM = '#,##0.00;(#,##0.00);"-"'


def _sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill = HDR, HFILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "B2"
    return ws


def write_report(state: SplitState) -> dict:
    results: list[MonthResult] = state["results"]
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    lines = [
        (f"Monthly split review - FY{state['fy'].fy_end_year} ({state['fy'].start} to {state['fy'].end})", True),
        (f"Source: {os.path.basename(state['workbook'])}, read as is. Nothing in QuickBooks or the workbook was changed.", False),
        ("", False),
        ("Cash test - B6 ('visa payment due') is booked in QB as PAID on 22200 (cheque visa portion C6 + "
         "'hygiene supplies' journal B6 - C6). 'Hygiene paid' is what actually left Hygiene's bank to TD after that "
         "statement and before the next. Cash diff = recorded as paid but not paid.", False),
        ("Expense test - the split should apply to the card cycle's own charges: new balance - previous balance + "
         "real payments. Lab fees count at 100% only when the charge is on that statement. "
         "DPC diff = correct Dental visa portion - booked (C6); + means Dental was under-charged.", False),
        ("", False),
        ("Proposed adjustments are for your accountant to review - not entries made by this tool.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = readme.cell(row=i, column=1, value=t)
        c.font, c.alignment = (FB if b else F), Alignment(wrap_text=True)
    readme.column_dimensions["A"].width = 140

    cols = ["Sheet", "Cheque #", "Cheque date", "Split cheque (C8)", "Hygiene deposit", "Card statement",
            "Visa due (B6)", "Hygiene paid TD", "Cash diff", "Cycle charges", "Lab on card", "Booked DPC (C6)",
            "Correct DPC", "DPC diff", "QB cheque", "QB journal", "Status"]
    ws = _sheet(wb, "Summary", cols, [11, 9, 12, 13, 12, 12, 13, 13, 12, 13, 12, 13, 13, 12, 9, 9, 14])
    for i, r in enumerate(results, 2):
        m = r.month
        row = [m.sheet, r.cheque_number, r.cheque_date, round(m.cheque_total, 2), r.deposit_date, r.statement_date,
               m.visa_due, r.bank_paid, r.cash_diff, r.expense_base, r.lab_on_card, round(m.visa_dpc, 2),
               r.correct_dpc, r.dpc_diff,
               None if r.qb_cheque_found is None else ("yes" if r.qb_cheque_found else "NO"),
               None if r.qb_journal_found is None else ("yes" if r.qb_journal_found else "NO"), r.status]
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = F
            if 4 <= j <= 14 and j not in (5, 6):
                c.number_format = NUM
            if j in (3, 5, 6):
                c.number_format = "yyyy-mm-dd"
        if r.status not in ("OK",):
            ws.cell(row=i, column=17).fill = ISSUE
    t = len(results) + 2
    ws.cell(row=t, column=1, value="Total").font = FB
    for j in (4, 7, 8, 9, 10, 11, 12, 13, 14):
        L = get_column_letter(j)
        c = ws.cell(row=t, column=j, value=f"=SUM({L}2:{L}{t - 1})")
        c.font, c.number_format = FB, NUM
    if state.get("boundary_results"):
        ws.cell(row=t + 2, column=1, value="Boundary months - linked and checked, NOT in the totals above").font = FB
        for i, r in enumerate(state["boundary_results"], t + 3):
            m = r.month
            row = [m.sheet, r.cheque_number, r.cheque_date, round(m.cheque_total, 2), r.deposit_date, r.statement_date,
                   m.visa_due, r.bank_paid, r.cash_diff, r.expense_base, r.lab_on_card, round(m.visa_dpc, 2),
                   r.correct_dpc, r.dpc_diff,
                   None if r.qb_cheque_found is None else ("yes" if r.qb_cheque_found else "NO"),
                   None if r.qb_journal_found is None else ("yes" if r.qb_journal_found else "NO"), r.status]
            for j, v in enumerate(row, 1):
                c = ws.cell(row=i, column=j, value=v)
                c.font = Font(name="Arial", italic=True)
                if 4 <= j <= 14 and j not in (5, 6):
                    c.number_format = NUM
                if j in (3, 5, 6):
                    c.number_format = "yyyy-mm-dd"

    adj = _sheet(wb, "Proposed adjustments", ["Sheet", "Card statement", "Area", "Amount", "Explanation"],
                 [11, 13, 22, 13, 110])
    n = 2
    for r in results:
        for area, amt, why in [
            ("Dental visa expense", r.dpc_diff,
             f"correct Dental visa portion {r.correct_dpc:,.2f} vs booked {r.month.visa_dpc:,.2f}" if r.correct_dpc is not None else ""),
            ("22200 paid but not paid", r.cash_diff,
             f"B6 {r.month.visa_due:,.2f} booked as paid; Hygiene paid {r.bank_paid:,.2f}" if r.bank_paid is not None else "")]:
            if amt is not None and abs(amt) >= 0.01:
                for j, v in enumerate([r.month.sheet, r.statement_date, area, amt, why], 1):
                    c = adj.cell(row=n, column=j, value=v)
                    c.font = F
                    if j == 4: c.number_format = NUM
                    if j == 2: c.number_format = "yyyy-mm-dd"
                n += 1
    adj.cell(row=n + 1, column=3, value="Total Dental visa expense").font = FB
    c = adj.cell(row=n + 1, column=4, value=f'=SUMIF(C2:C{n},"Dental visa expense",D2:D{n})')
    c.font, c.number_format = FB, NUM
    adj.cell(row=n + 2, column=3, value="Total 22200 paid but not paid").font = FB
    c = adj.cell(row=n + 2, column=4, value=f'=SUMIF(C2:C{n},"22200 paid but not paid",D2:D{n})')
    c.font, c.number_format = FB, NUM

    lab = _sheet(wb, "Lab fees", ["Sheet", "Description (workbook)", "Amount", "Found on card", "Card date"],
                 [11, 50, 12, 13, 12])
    n = 2
    for r in results:
        for item in r.month.lab_items:
            hit = r.lab_matches.get(item.row)
            for j, v in enumerate([r.month.sheet, item.description, item.amount, "yes" if hit else "NO",
                                   hit[1] if hit else None], 1):
                c = lab.cell(row=n, column=j, value=v)
                c.font = F
                if j == 3: c.number_format = NUM
                if j == 5: c.number_format = "yyyy-mm-dd"
            if not hit:
                lab.cell(row=n, column=4).fill = ISSUE
            n += 1

    cats = ["salary", "rent", "lease", "bookkeeping", "visa"]
    qbp = _sheet(wb, "QB postings", ["Sheet", "Cheque #"] + [f"{c} -> Dental account" for c in cats],
                 [11, 9] + [34] * len(cats))
    for i, r in enumerate(results, 2):
        qbp.cell(row=i, column=1, value=r.month.sheet).font = F
        qbp.cell(row=i, column=2, value=r.cheque_number).font = F
        for j, cat in enumerate(cats, 3):
            p = r.qb_postings.get(cat)
            qbp.cell(row=i, column=j, value=f"{p[0]}  ({p[1]:,.2f})" if p else None).font = F

    rf = state.get("rollforward")
    if rf is not None:
        rs = _sheet(wb, "22200 roll-forward", ["", f"{rf.open_date}", f"{rf.close_date}"], [70, 16, 16])
        rs.cell(row=1, column=1, value="Card account: QuickBooks 22200 vs the real TD Visa balance").font = HDR
        rows = [("QB 22200 - amount owed per QuickBooks", rf.qb_open, rf.qb_close),
                ("TD Visa - amount actually owed (statements)", rf.card_open, rf.card_close)]
        for i, (label, a, b) in enumerate(rows, 2):
            rs.cell(row=i, column=1, value=label).font = F
            for j, v in ((2, a), (3, b)):
                c = rs.cell(row=i, column=j, value=v); c.font, c.number_format = F, NUM
        rs.cell(row=4, column=1, value="Difference (QB - card)").font = FB
        for j, L in ((2, "B"), (3, "C")):
            c = rs.cell(row=4, column=j, value=f"=IF(OR({L}2=\"\",{L}3=\"\"),\"\",{L}2-{L}3)"); c.font, c.number_format = FB, NUM
        rs.cell(row=6, column=1, value="How the difference moved during the year").font = FB
        rs.cell(row=7, column=1, value="Opening difference").font = F
        c = rs.cell(row=7, column=3, value="=B4"); c.font, c.number_format = F, NUM
        n = 8
        for g, total, cnt in rf.groups():
            rs.cell(row=n, column=1, value=f"{g}  ({cnt} lines)").font = F
            c = rs.cell(row=n, column=3, value=total); c.font, c.number_format = F, NUM
            n += 1
        rs.cell(row=n, column=1, value="Unexplained").font = FB
        c = rs.cell(row=n, column=3, value=f"=IF(OR(C4=\"\",C7=\"\"),\"\",C4-C7-SUM(C8:C{n - 1}))"); c.font, c.number_format = FB, NUM
        rs.cell(row=n + 1, column=1, value="Closing difference").font = FB
        c = rs.cell(row=n + 1, column=3, value=f"=IF(C7=\"\",\"\",C7+SUM(C8:C{n}))"); c.font, c.number_format = FB, NUM
        for k, note in enumerate(rf.notes, n + 3):
            rs.cell(row=k, column=1, value="Note: " + note).font = Font(name="Arial", italic=True, color="C00000")
        rs.cell(row=n + 3 + len(rf.notes), column=1,
                value="+ = QuickBooks shows MORE owed on the card than TD does; - = less. Detail: 'Roll-forward lines' sheet.").font = F

        rd = _sheet(wb, "Roll-forward lines", ["Group", "Side", "Date", "Description", "Amount (as recorded)",
                                               "Effect on difference", "Note"], [58, 6, 11, 52, 14, 14, 70])
        for i, it in enumerate(rf.items, 2):
            for j, v in enumerate([it.group, it.side, it.txn_date, it.description, it.amount, it.effect, it.note], 1):
                c = rd.cell(row=i, column=j, value=v); c.font = F
                if j in (5, 6): c.number_format = NUM
                if j == 3: c.number_format = "yyyy-mm-dd"
            if it.note:
                rd.cell(row=i, column=7).fill = ISSUE

    if state.get("cutoffs"):
        co_ws = _sheet(wb, "Cut-off", ["", "Card statement", "Charges from", "Charges to", "Charges", "Lab fees (100% Dental)",
                                       "Hygiene share (20% of the rest)"], [60, 14, 12, 12, 13, 14, 16])
        for i, co in enumerate(state["cutoffs"], 2):
            vals = [co.label, co.statement_date, co.window[0], co.window[1], co.charges, co.lab, None]
            for j, v in enumerate(vals, 1):
                c = co_ws.cell(row=i, column=j, value=v); c.font = F
                if j in (2, 3, 4): c.number_format = "yyyy-mm-dd"
                if j in (5, 6): c.number_format = NUM
            c = co_ws.cell(row=i, column=7, value=f"=ROUND((E{i}-F{i})*0.2,2)"); c.font, c.number_format = FB, NUM
        k = len(state["cutoffs"]) + 3
        co_ws.cell(row=k, column=1, value="Card charges are booked 100% in Dental when they happen; Hygiene's share comes off "
                                          "with the following month's split. 'Closing' = this year's charges whose split "
                                          "falls in NEXT year; 'Opening' = last year's charges split in THIS year.").font = F
        co_ws.cell(row=k + 1, column=1, value="Lab fees are recognised by the lab vendors listed in the split workbook.").font = F

    fnd = _sheet(wb, "Findings", ["Sheet", "Area", "Severity", "Amount", "Message"], [11, 10, 9, 12, 120])
    order = {"Issue": 0, "Warning": 1, "Info": 2}
    for i, f in enumerate(sorted(state["findings"], key=lambda f: (order[f["severity"]], f["subject"])), 2):
        for j, v in enumerate([f["subject"].replace("Split ", ""), f["area"], f["severity"], f.get("amount"),
                               f["message"]], 1):
            c = fnd.cell(row=i, column=j, value=v)
            c.font = F
            if j == 4: c.number_format = NUM
            if j == 5: c.alignment = Alignment(wrap_text=True, vertical="top")
        if f["severity"] == "Issue":
            fnd.cell(row=i, column=3).fill = ISSUE

    try:
        wb.save(state["report"])
        print(f"📊 write_report: {state['report']}")
    except PermissionError:
        print(f"❌ write_report: {state['report']} is open in Excel - close it and rerun")
    return {}


def stop(state: SplitState) -> dict:
    print(f"❌ {state['error']}")
    return {}


def build_graph():
    g = StateGraph(SplitState)
    for name, fn in [("read_workbook", read_workbook), ("load_evidence", load_evidence), ("assess", assess),
                     ("rollforward", rollforward), ("save_to_sql", save_to_sql), ("write_report", write_report),
                     ("stop", stop)]:
        g.add_node(name, fn)
    g.set_entry_point("read_workbook")
    g.add_conditional_edges("read_workbook", lambda s: "stop" if s["error"] else "load_evidence",
                            {"stop": "stop", "load_evidence": "load_evidence"})
    g.add_edge("load_evidence", "assess")
    g.add_edge("assess", "rollforward")
    g.add_edge("rollforward", "save_to_sql")
    g.add_edge("save_to_sql", "write_report")
    g.add_edge("write_report", END)
    g.add_edge("stop", END)
    return g.compile()


def main():
    global ENGINE
    ap = argparse.ArgumentParser(description="Check the monthly Dental / Hygiene split workbook against the bank, card and QB data.")
    ap.add_argument("--fy-end", type=int, default=FISCAL_YEAR_END)
    ap.add_argument("--workbook", help="Path to the split workbook (default: <year root>\\NM_<year>.xlsx)")
    ap.add_argument("--report", help="Where to write the review workbook")
    ap.add_argument("--root", help="Override the year folder")
    ap.add_argument("--show-graph", action="store_true")
    args = ap.parse_args()
    app = build_graph()
    if args.show_graph:
        print(app.get_graph().draw_mermaid())
        return
    root = args.root or YEAR_ROOT.format(year=args.fy_end)
    workbook = args.workbook or os.path.join(root, WORKBOOK.format(year=args.fy_end))
    report = args.report or os.path.join(root, REPORT.format(year=args.fy_end))
    if not os.path.exists(workbook):
        sys.exit(f"❌ Split workbook not found: {workbook}")
    ENGINE = get_engine()
    ensure_schema(ENGINE)
    fy = fiscal_period(args.fy_end)
    print(f"🧾 Split review FY{fy.fy_end_year}: {fy.start} .. {fy.end}")
    app.invoke({"workbook": workbook, "report": report, "fy": fy, "months": [], "boundary": [], "evidence": None,
                "results": [], "boundary_results": [], "rollforward": None, "cutoffs": [], "findings": [], "error": None})


if __name__ == "__main__":
    main()
